"""Post-polish QA for the ICH manuscript."""
from __future__ import annotations

import importlib.util
import json
import re
import difflib
from collections import Counter
from pathlib import Path

import docx
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "work"
DOCX = ROOT / "1025Manuscript.docx"
NEW_SOURCE = WORK / "manuscript_content.py"
OLD_SOURCE = WORK / "manuscript_content_before_language_polish.py"
NUMBER_PATTERN = re.compile(r"(?<![A-Za-z])[-+]?\d(?:\d*[\d,]*)?(?:\.\d+)?%?")
CITE_PATTERN = re.compile(r"\{cite:([^}]+)\}")
FIGURE_REFERENCE = re.compile(r"\bFigure\s+(\d+)([A-Z])?\b")
TABLE_REFERENCE = re.compile(r"\bTable\s+(\d+)\b")
SENTENCE_SPLIT = re.compile(r"(?<=[.!?])\s+(?=[A-Z0-9])")


def load_module(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def text_blocks(module) -> list[tuple[str, str]]:
    blocks: list[tuple[str, str]] = [
        ("abstract", f"{label}: {text}") for label, text in module.ABSTRACT
    ]
    for block in module.BODY:
        if block[0] in {"para", "heading1", "heading2"}:
            blocks.append((block[0], block[1]))
    return blocks


def numeric_tokens(blocks: list[tuple[str, str]]) -> Counter[str]:
    return Counter(
        token.replace(",", "")
        for _, text in blocks
        for token in NUMBER_PATTERN.findall(text)
    )


def citation_order(blocks: list[tuple[str, str]]) -> dict[str, int]:
    order: dict[str, int] = {}
    for _, text in blocks:
        for match in CITE_PATTERN.finditer(text):
            for key in match.group(1).split(","):
                key = key.strip()
                if key and key not in order:
                    order[key] = len(order) + 1
    return order


def paragraph_runs(paragraph) -> list:
    return [run for run in paragraph.runs if run.text.strip()]


def run_font_names(paragraph) -> set[str]:
    names = set()
    for run in paragraph.runs:
        if not run.text.strip():
            continue
        names.add(run.font.name or "")
        fonts = run._element.find(qn("w:rPr"))
        if fonts is not None:
            fonts = fonts.find(qn("w:rFonts"))
        if fonts is not None:
            for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                names.add(fonts.get(qn(attribute)) or "")
    return names


def body_sequence(document) -> list[tuple[str, object]]:
    sequence: list[tuple[str, object]] = []
    paragraphs = iter(document.paragraphs)
    tables = iter(document.tables)
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            sequence.append(("p", next(paragraphs)))
        elif child.tag == qn("w:tbl"):
            sequence.append(("tbl", next(tables)))
    return sequence


def table_is_three_line(table) -> bool:
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        return False
    values = {}
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        values[edge] = element.get(qn("w:val")) if element is not None else None
    return values == {
        "top": "single",
        "left": "none",
        "bottom": "single",
        "right": "none",
        "insideH": "none",
        "insideV": "none",
    }


def main() -> None:
    new = load_module(NEW_SOURCE, "manuscript_content_new")
    old = load_module(OLD_SOURCE, "manuscript_content_old")
    new_blocks = text_blocks(new)
    old_blocks = text_blocks(old)
    new_numbers = numeric_tokens(new_blocks)
    old_numbers = numeric_tokens(old_blocks)
    old_order = citation_order(old_blocks)
    new_order = citation_order(new_blocks)
    reference_order = json.loads(
        (WORK / "citation_order_before.json").read_text(encoding="utf-8")
    )

    print("== source comparison ==")
    print(f"old numeric tokens={sum(old_numbers.values())}")
    print(f"new numeric tokens={sum(new_numbers.values())}")
    print(f"numeric token multiset equal={new_numbers == old_numbers}")
    if new_numbers != old_numbers:
        print("new only:", new_numbers - old_numbers)
        print("old only:", old_numbers - new_numbers)
    print(f"citation keys old={len(old_order)} new={len(new_order)}")
    print(f"citation order equal={new_order == old_order}")
    print(
        "citation order matches saved baseline="
        f"{list(new_order) == reference_order}"
    )

    unresolved = [
        (kind, text) for kind, text in new_blocks if "{cite:" in text
    ]
    print(f"unresolved content citations={len(unresolved)}")

    document = docx.Document(str(DOCX))
    paragraph_text = [paragraph.text.strip() for paragraph in document.paragraphs]
    text = "\n".join(paragraph_text)
    abstract_end = next(
        index for index, value in enumerate(paragraph_text) if value == "Abstract"
    )
    keywords_index = next(
        index for index, value in enumerate(paragraph_text) if value.startswith("Keywords:")
    )
    abstract_paragraphs = paragraph_text[abstract_end + 1 : keywords_index]
    abstract_words = sum(len(value.split()) for value in abstract_paragraphs)

    print("\n== docx structure ==")
    print(f"paragraphs={len(document.paragraphs)} tables={len(document.tables)} images={len(document.inline_shapes)}")
    print(f"abstract labels={[value.split(':', 1)[0] for value in abstract_paragraphs]}")
    print(f"abstract whitespace tokens={abstract_words}")
    print(f"keywords={document.paragraphs[keywords_index].text}")

    sequence = body_sequence(document)
    figure_caption_below = True
    table_caption_above = True
    figure_captions = []
    table_captions = []
    for index, (kind, item) in enumerate(sequence):
        if kind == "p":
            has_drawing = bool(item._element.findall(".//" + qn("w:drawing")))
            if has_drawing:
                following = sequence[index + 1][1] if index + 1 < len(sequence) else None
                if following is None or not re.match(r"^Figure\s+\d+\.", following.text):
                    figure_caption_below = False
                else:
                    figure_captions.append(following.text[:40])
            if re.match(r"^Figure\s+\d+\.", item.text):
                previous = sequence[index - 1][0] if index else None
                if previous != "p":
                    figure_caption_below = False
        elif kind == "tbl":
            previous = sequence[index - 1][1] if index else None
            if previous is None or not re.match(r"^Table\s+\d+\.", previous.text):
                table_caption_above = False
            else:
                table_captions.append(previous.text[:40])

    print(f"figure captions below images={figure_caption_below}")
    print(f"table captions above tables={table_caption_above}")
    print(f"table captions={len(table_captions)} three-line={all(table_is_three_line(table) for table in document.tables)}")

    print("\n== figure/table references ==")
    referenced_figures = sorted({int(match[0]) for match in FIGURE_REFERENCE.findall(text)})
    referenced_tables = sorted({int(match) for match in TABLE_REFERENCE.findall(text)})
    print(f"referenced figures={referenced_figures}")
    print(f"referenced tables={referenced_tables}")
    print(f"figure count={len(document.inline_shapes)} table count={len(document.tables)}")

    print("\n== formatting ==")
    normal = document.styles["Normal"]
    print(f"Normal={normal.font.name} {normal.font.size.pt if normal.font.size else None}pt line_spacing={normal.paragraph_format.line_spacing}")
    off_font = []
    off_body_size = []
    for index, paragraph in enumerate(document.paragraphs):
        if not paragraph.text.strip():
            continue
        fonts = run_font_names(paragraph)
        if fonts and fonts != {"Times New Roman"}:
            off_font.append((index, sorted(fonts), paragraph.text[:60]))
        sizes = {
            run.font.size.pt
            for run in paragraph.runs
            if run.text.strip() and run.font.size is not None
        }
        if sizes and sizes <= {18.0, 14.0, 12.0, 10.0} and 12.0 not in sizes and paragraph.style.name == "Normal":
            if not paragraph.text.startswith(("Table ", "Figure ")):
                off_body_size.append((index, sorted(sizes), paragraph.text[:60]))
    print(f"non-TNR paragraphs={len(off_font)}")
    if off_font:
        print(off_font[:5])
    print(f"normal-style paragraphs without 12pt run={len(off_body_size)}")
    if off_body_size:
        print(off_body_size[:5])

    print("\n== language checks ==")
    prose_paragraphs = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.style.name != "RefEntry"
        and paragraph.text.strip()
        and not paragraph.text.startswith(("Figure ", "Table "))
    ]
    prose_text = "\n".join(prose_paragraphs)
    british = re.findall(
        r"\b(?:haemorrhage\w*|regularis\w*|stabilis\w*|standardis\w*|"
        r"summaris\w*|normalis\w*|favour\w*|centre\w*|labelled|"
        r"haemodynamic\w*|regularised)\b",
        prose_text,
        flags=re.IGNORECASE,
    )
    print(f"British spelling hits outside references={british}")
    forbidden = [
        "511",
        "mRS",
        "simulated cohort",
        "simulation cohort",
        "local cohort",
        "clinical cohort",
        "hospital cohort",
    ]
    for phrase in forbidden:
        print(f"forbidden {phrase!r} count={text.count(phrase)}")

    body_text = prose_text
    sentences = [
        sentence.strip()
        for sentence in SENTENCE_SPLIT.split(body_text)
        if sentence.strip()
    ]
    lengths = [len(sentence.split()) for sentence in sentences]
    over_30 = [(length, sentence[:100]) for length, sentence in zip(lengths, sentences) if length > 30]
    print(f"sentences={len(sentences)} over30={len(over_30)} max={max(lengths)}")
    print("longest:")
    for length, sentence in sorted(over_30, reverse=True)[:5]:
        print(f"  {length}: {sentence}")

    repeated_words = re.findall(
        r"\b([A-Za-z]+)\s+\1\b",
        prose_text,
        flags=re.IGNORECASE,
    )
    print(f"immediate repeated words={repeated_words}")
    print(f"double spaces={prose_text.count('  ')}")

    (WORK / "manuscript_after_polish.txt").write_text(
        "\n\n".join(prose_paragraphs) + "\n",
        encoding="utf-8",
    )
    plain_lines = ["# Manuscript", "", "## Abstract", ""]
    plain_lines.extend(
        f"{label}: {text}" for label, text in new.ABSTRACT
    )
    plain_lines.extend(["", f"Keywords: {new.KEYWORDS}", ""])
    for kind, value in new.BODY:
        if kind == "heading1":
            plain_lines.extend([f"# {value}", ""])
        elif kind == "heading2":
            plain_lines.extend([f"## {value}", ""])
        elif kind == "para":
            plain_lines.extend([value, ""])
    (WORK / "manuscript_plain_after.txt").write_text(
        "\n".join(plain_lines) + "\n",
        encoding="utf-8",
    )
    old_text = "\n\n".join(text for _, text in old_blocks)
    new_text = "\n\n".join(text for _, text in new_blocks)
    (WORK / "language_polish_changes.diff").write_text(
        "\n".join(
            difflib.unified_diff(
                old_text.splitlines(),
                new_text.splitlines(),
                fromfile="before_language_polish",
                tofile="after_language_polish",
                lineterm="",
            )
        )
        + "\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
