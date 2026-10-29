"""Build a paragraph-level outline compliance matrix.

The matrix uses only four statuses: PASS, PARTIAL, UNSUPPORTED, and FAIL.
It records the current preformal state and can be rerun after the formal run
without changing the outline itself.
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document


ROOT = Path(__file__).resolve().parents[1]
OUTLINE_DOCX = ROOT / "outline.docx"
PROCESSED = ROOT / "data" / "processed"
COHORTS = ("mimic4", "mimic3", "eicu")
AUDIT_DIR = ROOT / "results" / "audits"
JSON_OUTPUT = AUDIT_DIR / "outline_compliance_matrix.json"
MARKDOWN_OUTPUT = AUDIT_DIR / "outline_compliance_matrix.md"

ALLOWED_STATUSES = {"PASS", "PARTIAL", "UNSUPPORTED", "FAIL"}


def _row(
    status: str,
    gate: str,
    evidence: str,
    action: str,
) -> dict[str, str]:
    if status not in ALLOWED_STATUSES:
        raise ValueError(f"Invalid compliance status: {status}")
    return {
        "status": status,
        "gate": gate,
        "evidence": evidence,
        "action": action,
    }


def _flow_count(flow: pd.DataFrame, reason_code: str) -> int:
    match = flow[flow["reason_code"].astype(str).eq(reason_code)]
    if len(match) != 1:
        raise ValueError(
            f"Cohort flow must contain one row for {reason_code!r}; "
            f"observed {len(match)}."
        )
    return int(match["n"].iloc[0])


def _flow_count_any(flow: pd.DataFrame, reason_codes: tuple[str, ...]) -> int:
    match = flow[flow["reason_code"].astype(str).isin(reason_codes)]
    if len(match) != 1:
        raise ValueError(
            f"Cohort flow must contain one row for one of {reason_codes!r}; "
            f"observed {len(match)}."
        )
    return int(match["n"].iloc[0])


def _cohort_facts() -> dict[str, dict[str, int | float]]:
    facts: dict[str, dict[str, int | float]] = {}
    for cohort in COHORTS:
        frame = pd.read_csv(PROCESSED / cohort / "cohort.csv")
        flow = pd.read_csv(PROCESSED / cohort / "cohort_flow.csv")
        outcome = pd.to_numeric(frame["outcome"], errors="raise")
        death_72h = frame["death_within_72h"].fillna(False).astype(bool)
        if "hospital_expire_flag" in frame:
            hospital_death = (
                pd.to_numeric(frame["hospital_expire_flag"], errors="coerce")
                .eq(1)
                .fillna(False)
            )
        else:
            hospital_death = (
                frame["hospitaldischargestatus"]
                .astype(str)
                .str.casefold()
                .eq("expired")
            )
        rass_observed = pd.to_numeric(
            frame["min_rass_24h"], errors="coerce"
        ).notna()
        facts[cohort] = {
            "source": _flow_count(flow, "source_ich_icu_stays"),
            "common_candidates": _flow_count(
                flow, "source_common_landmark_candidates"
            ),
            "primary": _flow_count_any(
                flow,
                ("primary_analytic_cohort", "primary_complete_followup"),
            ),
            "death_72h_pct": float(100.0 * death_72h.mean()),
            "hospital_mortality_pct": float(100.0 * hospital_death.mean()),
            "any_event_pct": float(100.0 * outcome.ge(1).mean()),
            "rass_pct": float(100.0 * rass_observed.mean()),
        }
    return facts


def _joined(
    field: str,
    facts: dict[str, dict[str, int | float]],
    digits: int = 0,
) -> str:
    values = [facts[cohort][field] for cohort in COHORTS]
    if digits:
        return "/".join(f"{float(value):.{digits}f}" for value in values)
    return "/".join(str(int(value)) for value in values)


def _comma_percent(
    field: str,
    facts: dict[str, dict[str, int | float]],
    digits: int = 1,
) -> str:
    values = [facts[cohort][field] for cohort in COHORTS]
    return (
        f"{float(values[0]):.{digits}f}%, "
        f"{float(values[1]):.{digits}f}%, and "
        f"{float(values[2]):.{digits}f}%"
    )


def _flow_sequence(
    cohort: str,
    facts: dict[str, dict[str, int | float]],
) -> str:
    values = facts[cohort]
    return (
        f"{int(values['source']):,} -> "
        f"{int(values['common_candidates']):,} -> "
        f"{int(values['primary']):,}"
    )


def _render_fact_template(
    text: str,
    facts: dict[str, dict[str, int | float]],
) -> str:
    """Render dynamic cohort facts after the reporting artifacts are frozen."""
    return text.format(
        flow_mimic4=_flow_sequence("mimic4", facts),
        flow_mimic3=_flow_sequence("mimic3", facts),
        flow_eicu=_flow_sequence("eicu", facts),
        death_72h_pct=_comma_percent("death_72h_pct", facts),
        hospital_mortality_pct=_comma_percent(
            "hospital_mortality_pct",
            facts,
        ),
        any_event_pct=_comma_percent("any_event_pct", facts),
        rass_pct=_comma_percent("rass_pct", facts),
    )


# The mapping covers every non-empty paragraph in outline.docx. Structural
# headings are included so that a missing or renumbered outline paragraph is
# detected rather than silently omitted.
CLAUSES: dict[str, dict[str, str]] = {
    "P000": _row("PASS", "structure", "Title heading is present.", "None."),
    "P001": _row(
        "PASS",
        "structure",
        "The title states the population, ordinal endpoint, dynamic graph network, and temporal regularization.",
        "None.",
    ),
    "P002": _row("PASS", "structure", "Highlights heading is present.", "None."),
    "P003": _row(
        "PASS",
        "design",
        "The code applies a 24 h predictor window, 48 h outcome horizon, ordinal endpoint, three databases, and direct transport.",
        "None.",
    ),
    "P004": _row("PASS", "structure", "Keywords heading is present.", "None."),
    "P005": _row(
        "PASS",
        "structure",
        "All five requested keywords are present.",
        "None.",
    ),
    "P006": _row("PASS", "structure", "Abstract heading is present.", "None."),
    "P007": _row(
        "PARTIAL",
        "formal_results",
        "The abstract generator is wired to formal metrics, but the outline still contains result placeholders and the formal run has not completed.",
        "Populate only from a validated formal manifest; retain neutral wording if TR has no discrimination gain.",
    ),
    "P008": _row("PASS", "structure", "Introduction heading is present.", "None."),
    "P009": _row(
        "PARTIAL",
        "reporting",
        "The outline mortality ranges are not assumed to match the rebuilt "
        "cohort. Observed 72 h mortality is {death_72h_pct}; in-hospital "
        "mortality is {hospital_mortality_pct}.",
        "Replace the draft ranges in generated prose with cohort-specific audited counts and denominators.",
    ),
    "P010": _row(
        "PASS",
        "design",
        "The introduction's clinical rationale and GCS-centered endpoint framing are consistent with the analysis.",
        "None.",
    ),
    "P011": _row(
        "PASS",
        "design",
        "Static-score and single-center limitations motivate the dynamic, multi-database design.",
        "None.",
    ),
    "P012": _row(
        "PASS",
        "design",
        "Dynamic change, ordinal stratification, and documentation heterogeneity are explicit design targets.",
        "None.",
    ),
    "P013": _row(
        "PASS",
        "design",
        "The design includes a common operational endpoint and prespecified definition sensitivity.",
        "None.",
    ),
    "P014": _row("PASS", "structure", "Methods heading is present.", "None."),
    "P015": _row("PASS", "structure", "Data sources subsection is present.", "None."),
    "P016": _row(
        "PASS",
        "data",
        "MIMIC-IV v3.1, MIMIC-III v1.4, and eICU-CRD v2.0 are bound in source and reporting code; GCS documentation differences are documented.",
        "None.",
    ),
    "P017": _row("PASS", "structure", "Study cohort subsection is present.", "None."),
    "P018": _row(
        "PARTIAL",
        "data",
        "The rebuilt cohort artifacts implement unique flows for MIMIC-IV "
        "({flow_mimic4}), MIMIC-III ({flow_mimic3}), and eICU "
        "({flow_eicu}). Every included patient, including a death by hour 72, "
        "must have at least two distinct ICU hours with valid "
        "total-GCS values between ICU hours 24 and 72; sparse-GCS deaths are "
        "excluded. A live ICU discharge before hour 72 is retained when the "
        "observed window yields a classifiable grade, and the "
        "complete-follow-up subset is a prespecified sensitivity analysis. "
        "The outline's draft 3,688/1,795/3,601 and "
        "2,851/1,322/2,032 counts differ from the audited flow and are stale.",
        "Read source and primary counts directly from rebuilt cohort_flow.csv "
        "and cohort.csv. Do not repeat the outline's stale sample sizes or "
        "waive the post-landmark GCS requirement for death.",
    ),
    "P019": _row("PASS", "structure", "Outcome definition subsection is present.", "None."),
    "P020": _row(
        "PARTIAL",
        "endpoint",
        "P020 = PARTIAL and p020_complete = false. Implemented: 24 h "
        "maximum-GCS baseline, two adjacent hourly-median observations at "
        "least 2 points below baseline, a 24-72 h window, and death within "
        "72 h of ICU admission as Grade 2 only after the strict P018 eligibility rule, "
        "including at least two distinct valid post-landmark GCS hours, has "
        "been met. Sparse-GCS deaths are excluded. Not implemented: verified "
        "post-landmark unplanned intubation, because planned status is "
        "unavailable.",
        "Keep P020 labelled PARTIAL. Do not proxy unplanned intubation with "
        "ordinary procedure records, and do not waive P018 for death.",
    ),
    "P021": _row("PASS", "structure", "Predictive variables subsection is present.", "None."),
    "P022": _row(
        "PARTIAL",
        "data",
        "All 14 configured channels exist in every tensor; gcs_total is the "
        "added 14th channel. Saved tensors preserve missingness. "
        "SequencePreprocessor is fitted within each outer training fold using "
        "training rows only, then applies causal within-stay forward fill "
        "(never backward fill). Observed first-24-hour RASS coverage is "
        "{rass_pct} in MIMIC-IV, MIMIC-III, and eICU; coverage is read from "
        "the rebuilt cohort artifacts. eICU urine output and FiO2 are "
        "completely missing.",
        "Report fold-fitted causal forward fill and channel coverage explicitly. Do not repeat the outline claim of complete RASS availability or claim observed eICU urine/FiO2 values.",
    ),
    "P023": _row("PASS", "structure", "Model subsection is present.", "None."),
    "P024": _row(
        "PASS",
        "model",
        "The five baselines, DG-GRU, two cumulative-logit branches, 18-value lambda grid, temporal penalty, and nested stacking are implemented. Four matched internal-development ablations isolate sequence versus static-summary information and the trajectory mask; they are not externally transported.",
        "Freeze runtime source before the formal run and report all matched ablations without externalizing them.",
    ),
    "P025": _row("PASS", "structure", "Training and validation subsection is present.", "None."),
    "P026": _row(
        "PASS",
        "protocol",
        "Patient-grouped 5-fold CV, 5 repeats, five fixed model seeds, and an explicit hyperparameter grid are implemented. Validation probabilities averaged over the five seeds select inner candidates and lambda. Primary paired inference is the 2,000-replicate conditional cluster bootstrap (patients in MIMIC; hospital_key in eICU). Standard DeLong is unclustered, first-CV-partition, and secondary. The P024/P039 fixed primary Holm family contains 10 tests for the five baseline models (xgb_static, xgb_summary, gru, grud, transformer) x any/severe; DG-GRU, GCS slope, stacking, and the three controls are exploratory_not_in_primary_holm_family.",
        "Execute exactly once on frozen source and data. Report hospital_key-cluster bootstrap as the eICU primary inference. Report DeLong only under the fixed 10-test family and mark every other comparator exploratory_not_in_primary_holm_family.",
    ),
    "P027": _row("PASS", "structure", "Statistical analysis subsection is present.", "None."),
    "P028": _row(
        "PASS",
        "metrics",
        "Macro AUC, per-grade AUC, adjacent-class error, Brier score, calibration, NRI, IDI, and decision-curve functions are implemented.",
        "Verify all reported metrics after the formal run.",
    ),
    "P029": _row("PASS", "structure", "Sensitivity analyses subsection is present.", "None."),
    "P030": _row(
        "PARTIAL",
        "sensitivity",
        "Eight analysis families are registered. The unplanned-intubation composite is unsupported in all datasets; sedation scenarios are unsupported in MIMIC-III/eICU; RASS-adjusted eICU GCS is unsupported.",
        "Report unsupported cells explicitly in Table 4. Do not substitute ordinary intubation, missing RASS as no sedation, or an ad hoc adjusted GCS.",
    ),
    "P031": _row("PASS", "structure", "Ethics and data availability subsection is present.", "None."),
    "P032": _row(
        "PASS",
        "reporting",
        "The manuscript and release tooling separate aggregate/open code from licensed patient-level data.",
        "Preserve the data-use restrictions in the final release.",
    ),
    "P033": _row("PASS", "structure", "Results heading is present.", "None."),
    "P034": _row(
        "PARTIAL",
        "results",
        "Cohort, missingness, documentation-frequency, and flow builders exist but current results are pending formal completion and rendering.",
        "Generate Table 1 and Figure 1A only from audited cohort artifacts.",
    ),
    "P035": _row(
        "PARTIAL",
        "results",
        "The required clinical characteristics, missingness, and GCS documentation summaries can be generated.",
        "Verify every reported count against the final cohort audit.",
    ),
    "P036": _row(
        "PARTIAL",
        "results",
        "The three-grade outcome distribution and event-rate table can be generated.",
        "Populate from the final cohorts, not the stale outline rates.",
    ),
    "P037": _row(
        "PARTIAL",
        "reporting",
        "The outline's 43.4% MIMIC-III and 28.5% eICU rates are stale. "
        "Rebuilt any-event rates are {any_event_pct} for MIMIC-IV, "
        "MIMIC-III, and eICU. Documentation and calendar-era patterns can be "
        "described, but causal attribution is unsupported.",
        "Replace both rates and use descriptive, non-causal wording for the between-database difference.",
    ),
    "P038": _row(
        "PARTIAL",
        "formal_results",
        "Macro-AUC, per-grade AUC, paired intervals, and the GCS-slope-only comparator are implemented; formal estimates are pending.",
        "Report all prespecified comparisons without selecting a favorable subset.",
    ),
    "P039": _row(
        "PARTIAL",
        "formal_results",
        "The GCS-slope-only benchmark is registered and participates in the same evaluation path.",
        "Report it alongside DG-GRU + TR after the frozen formal run.",
    ),
    "P040": _row(
        "PARTIAL",
        "formal_results",
        "Grade-specific and cumulative calibration and decision-curve code exists; formal curves are pending.",
        "Do not call calibration maintained unless the formal slopes, intercepts, and curves support it.",
    ),
    "P041": _row(
        "PARTIAL",
        "formal_results",
        "Reliability, calibration, and 10%-40% decision-curve rendering is implemented.",
        "Regenerate from the validated formal predictions.",
    ),
    "P042": _row(
        "PARTIAL",
        "formal_results",
        "Landmark performance, trajectory smoothness, and regularization comparisons are implemented but not yet executed at all three landmarks.",
        "Run the prespecified 6/24/48 h sensitivity analyses and report all results.",
    ),
    "P043": _row(
        "PARTIAL",
        "sensitivity",
        "Time-reversed and time-permuted controls are implemented. The third "
        "control adds generic Laplace input noise with scale 0.01 only to "
        "observed standardized time-series values and applies no temporal "
        "smoothness penalty.",
        "Describe the third control as a generic Laplace input-noise "
        "comparator, not as a temporal-smoothness penalty arm.",
    ),
    "P044": _row(
        "PARTIAL",
        "formal_results",
        "Direct external transport to MIMIC-III and eICU is implemented without external refitting.",
        "Report external performance only after lineage validation.",
    ),
    "P045": _row(
        "PARTIAL",
        "sensitivity",
        "GCS documentation-intensity transport strata are implemented with fixed development cut points.",
        "Report all three strata externally; do not claim calendar causality.",
    ),
    "P046": _row(
        "PARTIAL",
        "formal_results",
        "Top-feature explanation, time-dependent importance, and five subgroup analyses are implemented.",
        "Bind explanations to the frozen formal model artifacts.",
    ),
    "P047": _row(
        "PARTIAL",
        "formal_results",
        "SHAP, time-dependent gradient importance, and age/ventilation/anticoagulation/neurosurgery/sedation subgroup code exists.",
        "Regenerate from the frozen run and label exploratory analyses as exploratory.",
    ),
    "P048": _row(
        "PARTIAL",
        "sensitivity",
        "Eight sensitivity families have executable or explicit unsupported status records.",
        "Render every family in one table, including unsupported cells.",
    ),
    "P049": _row(
        "PARTIAL",
        "sensitivity",
        "The consolidated sensitivity-table builder is present but final scenario outputs are pending.",
        "Report event-rate and macro-AUC changes where estimable; explain unsupported analyses.",
    ),
    "P050": _row("PASS", "structure", "Discussion heading is present.", "None."),
    "P051": _row(
        "PARTIAL",
        "formal_results",
        "Generated discussion wording is conditional on the formal estimates.",
        "Retain the claim only if discrimination and calibration meet the prespecified wording threshold.",
    ),
    "P052": _row(
        "PASS",
        "design",
        "The discussion's comparison with static and single-center approaches follows from the study design.",
        "None.",
    ),
    "P053": _row(
        "PASS",
        "design",
        "Definition sensitivity and reproducibility are explicit contributions.",
        "None.",
    ),
    "P054": _row(
        "PARTIAL",
        "reporting",
        "The design can quantify documentation-frequency strata and describe era differences, but cannot isolate calendar era or monitoring intensity as causal effects.",
        "Use association language and retain case-mix and measurement limitations.",
    ),
    "P055": _row(
        "PARTIAL",
        "formal_results",
        "Title and abstract switching logic is implemented: TR emphasis is retained only when the development paired interval supports a gain.",
        "Apply the switch from final paired estimates; do not tune to obtain a positive result.",
    ),
    "P056": _row(
        "PARTIAL",
        "reporting",
        "Core limitations are represented. Required additions are eICU "
        "urine/FiO2 non-observation, observed RASS coverage of {rass_pct}, "
        "P020 = PARTIAL, the strict P018 requirement of at least two distinct "
        "post-landmark GCS hours even for deaths, and exclusion of sparse-GCS "
        "deaths.",
        "Add these limitations and avoid claims of complete endpoint "
        "harmonization or complete RASS availability.",
    ),
    "P057": _row("PASS", "structure", "Conclusion heading is present.", "None."),
    "P058": _row(
        "PARTIAL",
        "formal_results",
        "The conclusion generator can report transport performance but must not assume consistency.",
        "Retain 'consistent discrimination and calibration' only if the formal external estimates support both.",
    ),
    "P059": _row(
        "PARTIAL",
        "reporting",
        "The analysis uses de-identified public databases under their data use agreements and does not redistribute records.",
        "Insert the public repository URL and the per-database access statements once the repository exists.",
    ),
    "P060": _row("PASS", "structure", "Results heading is present.", "None."),
    "P061": _row(
        "PARTIAL",
        "formal_results",
        "Cohort, outcome-distribution, discrimination, calibration, temporal, transport, explanation, and sensitivity builders exist.",
        "Populate every results section from the validated formal manifest.",
    ),
    "P062": _row("PASS", "structure", "Discussion heading is present.", "None."),
    "P063": _row(
        "PARTIAL",
        "formal_results",
        "The principal-findings paragraph is generated from formal estimates.",
        "Do not state internal or external discrimination until the formal run completes.",
    ),
    "P064": _row(
        "PARTIAL",
        "reporting",
        "The literature framing is drafted around single-centre END nomograms and radiomics models.",
        "Verify each cited comparator against the final reference list.",
    ),
    "P065": _row(
        "PASS",
        "sensitivity",
        "Prespecified baseline-definition scenarios (0-24 h maximum and final pre-landmark value) are registered.",
        "Report the event-rate shift produced by each definition in Table 4.",
    ),
    "P066": _row(
        "PARTIAL",
        "formal_results",
        "Transport strata by documentation intensity are implemented but not yet estimated.",
        "Report the cross-database difference descriptively; do not attribute it to case mix.",
    ),
    "P067": _row(
        "PARTIAL",
        "formal_results",
        "Title and abstract wording switch on the paired regularization contrast.",
        "Retain the regularization claim only if the paired interval supports it.",
    ),
    "P068": _row(
        "PARTIAL",
        "reporting",
        "Limitations cover sedation, absent imaging and NIHSS, eICU channel gaps, and live transfer handling.",
        "Replace the placeholder eICU coverage figure with the rebuilt cohort value.",
    ),
    "P069": _row("PASS", "structure", "Conclusion heading is present.", "None."),
    "P070": _row(
        "PARTIAL",
        "formal_results",
        "The conclusion text is generated from the validated manifest.",
        "Keep the conclusion to the prespecified wording once estimates are final.",
    ),
    "P071": _row(
        "PASS",
        "structure",
        "The figure and table registry is present.",
        "Regenerate the registry from the final figure outputs.",
    ),
}


def parse_outline() -> dict[str, str]:
    if not OUTLINE_DOCX.is_file():
        raise FileNotFoundError(f"Outline document not found: {OUTLINE_DOCX}")
    document = Document(OUTLINE_DOCX)
    texts = [
        " ".join(paragraph.text.split())
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]
    paragraphs = {
        f"P{index:03d}": text for index, text in enumerate(texts)
    }
    if not paragraphs:
        raise ValueError("No outline paragraphs were parsed")
    return paragraphs


def build_report() -> dict[str, Any]:
    paragraphs = parse_outline()
    missing = sorted(set(paragraphs) - set(CLAUSES))
    extra = sorted(set(CLAUSES) - set(paragraphs))
    if missing or extra:
        raise ValueError(
            f"Compliance coverage mismatch: missing={missing}, extra={extra}"
        )
    facts = _cohort_facts()
    rows = []
    for paragraph_id, text in paragraphs.items():
        clause = {
            key: _render_fact_template(value, facts)
            for key, value in CLAUSES[paragraph_id].items()
        }
        rows.append(
            {
                "paragraph_id": paragraph_id,
                "outline_text": text,
                **clause,
            }
        )
    counts = {
        status: sum(row["status"] == status for row in rows)
        for status in ("PASS", "PARTIAL", "UNSUPPORTED", "FAIL")
    }
    hard_failures = [row["paragraph_id"] for row in rows if row["status"] == "FAIL"]
    material_partials = [
        row["paragraph_id"]
        for row in rows
        if row["status"] in {"PARTIAL", "UNSUPPORTED"}
        and row["gate"] in {"endpoint", "data", "sensitivity", "formal_results"}
    ]
    return {
        "outline": str(OUTLINE_DOCX),
        "outline_source_format": "docx",
        "status_vocabulary": sorted(ALLOWED_STATUSES),
        "paragraph_coverage": {
            "outline_paragraphs": len(paragraphs),
            "mapped_paragraphs": len(rows),
            "complete": True,
        },
        "status_counts": counts,
        "hard_failures": hard_failures,
        "material_partials_or_unsupported": material_partials,
        "run_state": (
            "formal_result_pending"
            if not (ROOT / "results" / "audits" / "formal_run_validation.json").is_file()
            else "formal_validation_artifact_present"
        ),
        "rows": rows,
    }


def write_markdown(report: dict[str, Any]) -> None:
    lines = [
        "# Outline-to-Code Compliance Matrix",
        "",
        "This is the current paragraph-level audit of `outline.docx` against",
        "the implemented study pipeline. Status vocabulary is restricted to",
        "`PASS`, `PARTIAL`, `UNSUPPORTED`, and `FAIL`.",
        "",
        "## Summary",
        "",
        f"- Outline paragraphs mapped: {report['paragraph_coverage']['mapped_paragraphs']}",
        f"- PASS: {report['status_counts']['PASS']}",
        f"- PARTIAL: {report['status_counts']['PARTIAL']}",
        f"- UNSUPPORTED: {report['status_counts']['UNSUPPORTED']}",
        f"- FAIL: {report['status_counts']['FAIL']}",
        f"- Run state: `{report['run_state']}`",
        "",
        "The formal numerical results are intentionally not treated as complete",
        "until a validated frozen-source manifest and prediction artifacts exist.",
        "",
        "## Material deviations",
        "",
        "| Paragraph | Status | Gate | Finding | Required action |",
        "|---|---|---|---|---|",
    ]
    for row in report["rows"]:
        if row["paragraph_id"] in set(report["material_partials_or_unsupported"]):
            lines.append(
                "| {paragraph_id} | {status} | {gate} | {evidence} | {action} |".format(
                    **row
                )
            )
    lines.extend(
        [
            "",
            "## Full matrix",
            "",
            "| Paragraph | Status | Gate | Evidence | Action |",
            "|---|---|---|---|---|",
        ]
    )
    for row in report["rows"]:
        lines.append(
            "| {paragraph_id} | {status} | {gate} | {evidence} | {action} |".format(
                **row
            )
        )
    lines.append("")
    MARKDOWN_OUTPUT.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    AUDIT_DIR.mkdir(parents=True, exist_ok=True)
    report = build_report()
    JSON_OUTPUT.write_text(
        json.dumps(report, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    write_markdown(report)
    print(
        json.dumps(
            {
                "paragraphs": report["paragraph_coverage"]["mapped_paragraphs"],
                "status_counts": report["status_counts"],
                "json": str(JSON_OUTPUT),
                "markdown": str(MARKDOWN_OUTPUT),
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
