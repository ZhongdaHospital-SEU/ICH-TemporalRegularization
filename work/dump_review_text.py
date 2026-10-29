"""Dump the built manuscript with paragraph indices, tables and captions."""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("1025Manuscript.docx")
document = Document(path)

print("### PARAGRAPHS")
for index, paragraph in enumerate(document.paragraphs):
    text = paragraph.text.strip()
    if text:
        print(f"[{index:03d}] {text}")

print("\n### TABLES")
for table_index, table in enumerate(document.tables, start=1):
    print(f"--- table {table_index} rows={len(table.rows)} cols={len(table.columns)}")
    for row in table.rows:
        print(" | ".join(cell.text.strip() for cell in row.cells))
