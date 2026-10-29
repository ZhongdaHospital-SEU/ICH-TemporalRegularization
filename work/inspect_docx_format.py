"""Dump the paragraph and style structure of a reference manuscript."""
from __future__ import annotations

import sys
from pathlib import Path

import docx


def fmt(paragraph) -> str:
    runs = []
    for run in paragraph.runs[:3]:
        size = run.font.size.pt if run.font.size else None
        runs.append((run.bold, size, run.font.name))
    return f"{paragraph.style.name} | {runs} | {paragraph.text[:100]}"


def main() -> None:
    path = Path(sys.argv[1])
    limit = int(sys.argv[2]) if len(sys.argv) > 2 else 80
    document = docx.Document(str(path))
    section = document.sections[0]
    print(
        f"sections={len(document.sections)} tables={len(document.tables)} "
        f"images={len(document.inline_shapes)}"
    )
    print(
        f"page={section.page_width}x{section.page_height} "
        f"margins L{section.left_margin} R{section.right_margin} "
        f"T{section.top_margin} B{section.bottom_margin}"
    )
    style = document.styles["Normal"]
    print(
        f"Normal font={style.font.name} size={style.font.size.pt if style.font.size else None} "
        f"line_spacing={style.paragraph_format.line_spacing}"
    )
    for index, paragraph in enumerate(document.paragraphs[:limit]):
        print(index, "|", fmt(paragraph))
    print("--- paragraphs with style != Normal ---")
    for index, paragraph in enumerate(document.paragraphs):
        if paragraph.style.name != "Normal":
            print(index, "|", fmt(paragraph))


if __name__ == "__main__":
    main()
