"""Inspect the built manuscript for REF fields and matching bookmarks."""
from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

from docx import Document


def main() -> None:
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("1025Manuscript.docx")
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")

    instructions = re.findall(r"<w:instrText[^>]*>\s*REF\s+(\S+)[^<]*</w:instrText>", xml)
    bookmarks = re.findall(r'<w:bookmarkStart[^>]*w:name="([^"]+)"', xml)
    starts = re.findall(r"<w:bookmarkStart", xml)
    ends = re.findall(r"<w:bookmarkEnd", xml)
    print(f"ref_fields={len(instructions)} unique_targets={len(set(instructions))}")
    print(f"bookmarks={len(bookmarks)} start={len(starts)} end={len(ends)}")
    missing = sorted({name for name in instructions if name not in bookmarks})
    print(f"fields_without_bookmark={missing}")
    unreferenced = sorted({name for name in bookmarks if name not in instructions})
    print(f"bookmarks_without_field={unreferenced}")

    document = Document(path)
    bracket_groups = 0
    bad_groups = []
    for paragraph in document.paragraphs:
        text = paragraph.text
        for match in re.finditer(r"\[[^\]]*\]", text):
            group = match.group(0)
            if re.fullmatch(r"\[\d+(,\d+)*\]", group):
                bracket_groups += 1
            elif re.fullmatch(r"\[\s*\]", group):
                bad_groups.append(group)
    print(f"bracket_groups={bracket_groups} empty_groups={len(bad_groups)}")

    references = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name == "RefEntry"
    ]
    without_doi = [entry for entry in references if "https://doi.org/" not in entry]
    print(f"reference_entries={len(references)} without_doi={len(without_doi)}")
    for entry in without_doi:
        print("  ", entry[:120])


if __name__ == "__main__":
    main()
