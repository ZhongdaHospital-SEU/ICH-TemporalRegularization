from docx import Document

path = r"D:\TT paper\1024 TP Intracerebral Hemorrhage\1025Manuscript.docx"
doc = Document(path)

lines = []
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if not text:
        continue
    lines.append(f"[P{i:03d}] {text}")

lines.append("")
lines.append("==== TABLES ====")
for ti, table in enumerate(doc.tables, start=1):
    lines.append(f"--- TABLE OBJECT {ti} rows={len(table.rows)} cols={len(table.columns)} ---")
    for row in table.rows:
        cells = [c.text.replace("\n", " ").strip() for c in row.cells]
        lines.append(" | ".join(cells))

out = r"D:\TT paper\1024 TP Intracerebral Hemorrhage\work\review_dump_v2.txt"
with open(out, "w", encoding="utf-8") as fh:
    fh.write("\n".join(lines))
print("wrote", out, "lines", len(lines))
