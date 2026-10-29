"""Generate the TRIPOD+AI checklist in Markdown, CSV, and Word formats."""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from .common import OUTPUTS, RESULTS, load_config
from .make_manuscript import (
    OUTLINE_MODEL_ORDER,
    OUTLINE_REQUIRED_MODELS,
    lambda_frame,
    measurement_audit_frame,
    missing_figure_outputs,
    sensitivity_frame,
    shap_top15_frame,
    subgroup_frame,
    table3_comparison_frame,
    time_importance_frame,
    trajectory_frame,
)


TITLE = (
    "Temporal Regularization of a Dynamic Graph Network for Ordinal "
    "Prediction of Early Neurological Deterioration After Intracerebral "
    "Hemorrhage in Intensive Care"
)
DATABASE_SPANS = {
    "MIMIC-IV": "2008-2022",
    "MIMIC-III": "2001-2012",
    "eICU-CRD": "2014-2015",
}
ITEMS = [
    (
        "1",
        "Title identifies the study as developing and/or validating a "
        "prediction model, the target population, and the outcome.",
        "Title page",
        "Reported",
        "The title identifies ordinal prediction of early neurological "
        "deterioration or death after intracerebral hemorrhage in intensive "
        "care and names the temporal-regularization method.",
    ),
    (
        "2",
        "Abstract reports design, data sources, participants, outcome, "
        "predictors, sample size, performance, and limitations.",
        "Abstract",
        "Partially reported",
        "The finalized structured abstract will report the landmark, cohorts "
        "and outcome counts, predictors, cross-validation, external transport, "
        "principal discrimination and calibration results only after a "
        "validated formal run, together with the main limitations.",
    ),
    (
        "3",
        "Objectives and intended use are stated.",
        "Introduction; Discussion",
        "Reported",
        "The intended use is research-stage risk stratification at ICU hour "
        "24. The manuscript does not claim treatment selection, causal "
        "inference, or readiness for clinical deployment.",
    ),
    (
        "4",
        "Data sources, study design, and dates are described.",
        "Methods 2.1",
        "Reported",
        f"MIMIC-IV v3.1 ({DATABASE_SPANS['MIMIC-IV']}), MIMIC-III v1.4 "
        f"({DATABASE_SPANS['MIMIC-III']}), and eICU-CRD v2.0 "
        f"({DATABASE_SPANS['eICU-CRD']}) are identified. The design is a "
        "retrospective secondary analysis of de-identified critical care data.",
    ),
    (
        "5",
        "Setting and eligibility criteria, including centres, are described.",
        "Methods 2.1-2.2; Figure 1",
        "Reported",
        "Eligibility requires an adult first ICU episode with the prespecified "
        "ICH diagnosis codes (ICD-9 431/432; ICD-10 I61/I62), survival through "
        "ICU hour 24, at least one valid pre-landmark total GCS, and at least "
        "two distinct ICU hours with valid total-GCS values between ICU hours "
        "24 and 72. This post-landmark GCS requirement applies to every "
        "included patient, including those who died by hour 72; sparse-GCS "
        "deaths are excluded rather than assigned Grade 2. A live ICU "
        "discharge before hour 72 was not an exclusion criterion: the stay "
        "was retained when the observed window yielded a classifiable grade, "
        "and the complete-follow-up subset was analysed as a prespecified "
        "sensitivity analysis. The "
        "outline's earlier draft sample sizes differ from the rebuilt audit "
        "and are not repeated as current results. The code set does not "
        "establish strictly spontaneous etiology. MIMIC-III and eICU are "
        "external cohorts; MIMIC-III is single-institution and eICU is "
        "multi-hospital. The eICU hospital_key is retained for primary "
        "bootstrap uncertainty.",
    ),
    (
        "6",
        "Outcome definition, measurement, and timing are described.",
        "Methods 2.3; Supplementary Table S4",
        "Partially reported",
        "The common outcome is confirmed GCS deterioration or death from "
        "ICU hours 24 to 72. Grades are defined by a reduction of 2-3 points "
        "or at least 4 points. Death within 72 h of ICU admission, including "
        "a death after a live ICU discharge, is Grade 2 only after "
        "the strict P018 eligibility rule has been met, including at least two "
        "distinct valid post-landmark GCS hours; a sparse-GCS death is "
        "excluded. The endpoint is explicitly `P020 = PARTIAL` and "
        "`p020_complete = false` because procedure start times identify "
        "recorded intubation but not unplanned intubation; the required "
        "intubation-augmented analysis is therefore unsupported rather than "
        "being replaced by that proxy. The text states that GCS may be "
        "affected by sedation, analgesia, neuromuscular blockade, and "
        "ventilation.",
    ),
    (
        "7",
        "Predictors are defined, including measurement timing and availability "
        "at the landmark.",
        "Methods 2.4; Supplementary Table S2",
        "Reported",
        "All predictors end before ICU hour 24. Post-landmark observations, "
        "discharge location, and outcome-derived variables are excluded. "
        "Fourteen configured channels are represented by values and explicit "
        "observation masks. Total GCS reconciles the outline's stated 14-channel "
        "count with its 13 named modalities. Observed first-24-hour RASS "
        "coverage is read from the rebuilt cohort artifacts and reported per "
        "database; the outline's claim of full eICU RASS availability is not "
        "assumed. eICU-CRD also lacks structured urine output and FiO2 tables, "
        "so those masks remain missing rather than being imputed across "
        "databases.",
    ),
    (
        "8",
        "Sample size, number of events, and rationale are reported.",
        "Methods 2.2; Results 3.1",
        "Partially reported",
        "All eligible stays were used. Outcome counts and the flow from the "
        "initial ICH code query to the complete-follow-up cohort are reported "
        "for every database. No prospective sample-size calculation was "
        "performed; external estimates are explicitly labelled exploratory "
        "and are interpreted with their observed event counts and intervals.",
    ),
    (
        "9",
        "Missing-data handling is described and compatible with the prediction "
        "landmark.",
        "Methods 2.4; Supplementary Table S2",
        "Reported",
        "SequencePreprocessor is fitted within each outer training fold using "
        "training rows only, then applies causal within-stay forward filling "
        "across ICU hours 0-23; no backward filling is used. Static variables "
        "use training-fold medians and scaling. Channel-level observation masks "
        "preserve missingness information.",
    ),
    (
        "10",
        "Model type, algorithm, and rationale are described.",
        "Methods 2.5; Figure 2",
        "Reported",
        "The primary model combines a value-aware dynamic feature graph over "
        "14 value channels, a GRU, fusion of the 23-variable core static set, and two "
        "shared-hidden-state cumulative-logit branches for grade >=1 and "
        "grade 2. A temporal penalty regularizes the latent risk trajectory. "
        "The third negative control instead adds generic Laplace noise at "
        "scale 0.01 only to observed standardized temporal inputs and applies "
        "no temporal smoothness penalty. "
        "Summary XGBoost and the neural encoder use the same core static "
        "subset; the neural encoder "
        "additionally receives the hourly sequence, so cross-family results "
        "are pipeline comparisons.",
    ),
    (
        "11",
        "Model development, hyperparameter tuning, and internal validation are "
        "described.",
        "Methods 2.5-2.7",
        "Reported",
        "Repeated five-fold StratifiedGroupKFold separates patients. "
        "Preprocessing and early stopping are fitted inside each outer "
        "training fold. Five fixed seeds are used jointly in inner selection: "
        "validation probabilities are averaged across seeds before scoring "
        "candidate hyperparameters and lambda_monotone. Nested stacking uses "
        "five patient-grouped inner-fold OOF probabilities within each outer "
        "training split, then fits a multinomial logistic meta-model. External "
        "neural models use the median internal-CV epoch. Four matched "
        "internal-development ablations separate temporal inputs from static "
        "summaries; they are not externally transported.",
    ),
    (
        "12",
        "Model performance measures are defined.",
        "Methods 2.8",
        "Reported",
        "Primary discrimination is macro one-versus-rest AUC. Per-grade, "
        "cumulative any/severe AUC, adjacent and non-adjacent error, mean "
        "ordinal error, Brier score, log loss, calibration slope and intercept, "
        "NRI, IDI, ECE, and 10%-40% decision-curve analysis are reported.",
    ),
    (
        "13",
        "Model performance is reported with uncertainty.",
        "Results 3.2-3.6; Tables 2-5",
        "Partially reported",
        "After a validated formal run, primary conditional percentile bootstrap "
        "confidence intervals will use 2,000 paired resampling replicates, "
        "clustered by patient in MIMIC and by hospital_key in eICU. Models and "
        "realized CV partitions are held fixed, and paired AUC differences use "
        "the same cluster structure. Standard paired DeLong tests are "
        "unclustered, use the first CV partition only, and are secondary. The "
        "P024/P039 fixed primary Holm family contains 10 tests: the five baseline "
        "models (xgb_static, xgb_summary, gru, grud, transformer) x any/severe. DG-GRU, "
        "GCS slope, stacking, and the three controls are "
        "exploratory_not_in_primary_holm_family; the DeLong family does not "
        "replace the eICU hospital-cluster bootstrap for primary inference.",
    ),
    (
        "14",
        "Fairness and equity are considered, including subgroup performance.",
        "Results 3.6; Figure 6B; Supplementary Table S6",
        "Partially reported",
        "The validated formal report will describe development-cohort subgroup "
        "discrimination for age, mechanical ventilation, anticoagulation-related "
        "hemorrhage, neurosurgical intervention, and sedation exposure. "
        "Subgroup calibration, race-specific performance, and hospital-level "
        "calibration are exploratory and may remain unestimable with sparse "
        "events.",
    ),
    (
        "15",
        "Model interpretation is reported.",
        "Results 3.4 and 3.6; Figures 5 and 6A; Table 5; Tables S9-S10",
        "Partially reported",
        "After a validated formal run, TreeSHAP will rank the top 15 Summary "
        "XGBoost predictors, and time-dependent absolute-gradient importance "
        "will summarize DG-GRU + TR by channel and ICU hour. Cross-fitted "
        "latent trajectories will be compared with time-reversed, "
        "time-permuted, and generic Laplace input-noise controls. The Laplace "
        "arm perturbs only observed standardized temporal values at scale "
        "0.01 and applies no temporal smoothness penalty.",
    ),
    (
        "16",
        "Model output and its interpretation are described.",
        "Methods 2.5; Discussion 4.5",
        "Reported",
        "The output is a probability for each of three ordinal grades. Class "
        "labels use the median of the predicted ordinal distribution, with "
        "modal predictions reported secondarily. The latent risk score is not "
        "a validated biological quantity, and probabilities require local "
        "recalibration before absolute-risk use.",
    ),
    (
        "17",
        "External validation will be reported only from a validated run, with "
        "a clear statement of which model was transported.",
        "Methods 2.7; Results 3.2-3.4",
        "Partially reported",
        "The prespecified complete development pipeline, including "
        "preprocessing, hyperparameters, five seeds, and the nested stacking "
        "meta-model, will be transported without external refitting and "
        "reported only from a validated formal run. MIMIC-III is a "
        "same-institution source-shift cohort from a different database "
        "generation, and eICU is multi-hospital geographic transport.",
    ),
    (
        "18",
        "Model updating and recalibration are described.",
        "Discussion 4.4-4.5",
        "Partially reported",
        "External recalibration is not part of the primary analysis. A "
        "validated formal run will quantify the calibration gap. Local "
        "recalibration remains a required next step and will be kept separate "
        "from the external predictions used for evaluation.",
    ),
    (
        "19",
        "Real-world implementation, intended user, and workflow are discussed.",
        "Discussion 4.5",
        "Partially reported",
        "A plausible workflow is ICU risk review at hour 24 after medication "
        "and ventilation status are verified. No prospective silent trial, "
        "clinical impact study, or workflow evaluation was performed.",
    ),
    (
        "20",
        "Data availability is stated.",
        "Declarations",
        "Reported",
        "MIMIC-IV, MIMIC-III, and eICU-CRD are available from PhysioNet under "
        "credentialed access. Raw records are not redistributed. Derived "
        "aggregate results and reproducibility manifests will be supplied after "
        "the validated formal run.",
    ),
    (
        "21",
        "Code availability is stated.",
        "Declarations; README",
        "Reported",
        "The analysis code, configuration, tests, and reproduction commands "
        "are supplied with the manuscript package.",
    ),
    (
        "22",
        "Reporting-guideline adherence is stated.",
        "Methods; Declarations",
        "Reported",
        "The manuscript follows TRIPOD+AI and is accompanied by this checklist.",
    ),
    (
        "23",
        "Registration is stated.",
        "Declarations",
        "Not applicable",
        "The study is a retrospective analysis of existing databases and was "
        "not prospectively registered.",
    ),
    (
        "24",
        "Patient and public involvement is described.",
        "Declarations",
        "Not performed",
        "There was no patient or public involvement in the design, conduct, "
        "or reporting of this secondary analysis.",
    ),
    (
        "25",
        "Source of funding and role of funders are stated.",
        "Declarations",
        "To be completed",
        "No funding statement is asserted in this draft. The corresponding "
        "author must supply the funding source and funder role before "
        "submission.",
    ),
    (
        "26",
        "Conflicts of interest are declared.",
        "Declarations",
        "To be completed",
        "The author-level competing-interest declaration must be confirmed "
        "before submission.",
    ),
    (
        "27",
        "Use of artificial intelligence is disclosed.",
        "Declarations",
        "Reported",
        "An AI coding assistant supported code development, analysis "
        "documentation, and language editing. Analysis outputs are configured "
        "to be checked with unit tests and versioned reruns under fixed random "
        "seeds before reporting. "
        "Bitwise determinism is not claimed across hardware or library "
        "versions; no patient-level data were submitted to an external service.",
    ),
]


def write_csv() -> pd.DataFrame:
    frame = pd.DataFrame(
        ITEMS,
        columns=["Item", "Requirement", "Where addressed", "Status", "Comment"],
    )
    frame.to_csv(OUTPUTS / "TRIPOD_AI_checklist.csv", index=False, encoding="utf-8-sig")
    return frame


def write_markdown(frame: pd.DataFrame) -> None:
    lines = [
        "# TRIPOD+AI checklist",
        "",
        f"**Manuscript:** {TITLE}",
        "",
        "| Item | Requirement | Where addressed | Status | Comment |",
        "|---:|---|---|---|---|",
    ]
    for row in frame.itertuples(index=False):
        values = [
            str(row[0]),
            str(row[1]),
            str(row[2]),
            str(row[3]),
            str(row[4]),
        ]
        lines.append("| " + " | ".join(value.replace("|", "/") for value in values) + " |")
    (OUTPUTS / "TRIPOD_AI_checklist.md").write_text(
        "\n".join(lines) + "\n", encoding="utf-8"
    )


def set_cell_font(cell, size: float, bold: bool = False) -> None:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(size)
            run.bold = bold


def write_docx(frame: pd.DataFrame) -> None:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(8.5)

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run("TRIPOD+AI checklist")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(15)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(TITLE)
    run.italic = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)

    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    widths = [Cm(1.1), Cm(7.0), Cm(4.0), Cm(2.7), Cm(11.5)]
    headers = ["Item", "Requirement", "Where addressed", "Status", "Comment"]
    for cell, text, width in zip(table.rows[0].cells, headers, widths):
        cell.text = text
        cell.width = width
        set_cell_font(cell, 8.5, bold=True)
    for values in frame.itertuples(index=False):
        cells = table.add_row().cells
        for cell, text, width in zip(cells, values, widths):
            cell.text = str(text)
            cell.width = width
            set_cell_font(cell, 7.8)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(
        "TRIPOD+AI checklist for the ICH temporal-regularization study"
    )
    run.font.name = "Times New Roman"
    run.font.size = Pt(8)

    document.save(OUTPUTS / "TRIPOD_AI_checklist.docx")


def validate_reporting_dependency() -> None:
    """Refuse to certify a checklist that does not match the completed run."""
    manifest_path = RESULTS / "run_manifest.json"
    metrics_path = RESULTS / "metrics.json"
    if not manifest_path.exists() or not metrics_path.exists():
        raise FileNotFoundError(
            "TRIPOD+AI generation requires results/run_manifest.json and "
            "results/metrics.json."
        )
    config = load_config()
    if len(config.get("channels", [])) != 14:
        raise ValueError("configs/study.json must define 14 time-series channels.")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    declared_models = manifest.get("models", [])
    if isinstance(declared_models, str):
        declared_models = [declared_models]
    declared = {str(name) for name in declared_models}
    if declared != OUTLINE_REQUIRED_MODELS:
        missing = OUTLINE_REQUIRED_MODELS.difference(declared)
        unexpected = declared.difference(OUTLINE_REQUIRED_MODELS)
        raise ValueError(
            "The run manifest does not match the outline model set; "
            f"missing={sorted(missing)}, unexpected={sorted(unexpected)}."
        )
    seeds = manifest.get("model_seeds", [])
    if not isinstance(seeds, list) or len(seeds) != 5:
        raise ValueError("The outline requires five fixed model seeds.")
    if len({str(seed) for seed in seeds}) != 5:
        raise ValueError("The five fixed model seeds must be unique.")

    metrics = json.loads(metrics_path.read_text(encoding="utf-8"))
    # The checklist makes cohort-wide claims, so require the exact comparison
    # set and paired intervals in development, MIMIC-III, and eICU.
    table3_comparison_frame(metrics, OUTLINE_MODEL_ORDER)

    # Do not mark the associated items as reported unless the prespecified
    # sensitivity, trajectory, and explanation analyses satisfy their contracts.
    sensitivity_frame()
    trajectory_frame()
    subgroup_frame()
    shap_top15_frame()
    time_importance_frame(config)
    lambda_frame()
    measurement_audit_frame()
    missing_figures = missing_figure_outputs(RESULTS)
    if missing_figures:
        raise FileNotFoundError(
            "TRIPOD+AI generation requires every reported figure in both PNG "
            "and PDF form; missing: "
            + ", ".join(missing_figures)
        )


def main() -> None:
    validate_reporting_dependency()
    OUTPUTS.mkdir(parents=True, exist_ok=True)
    frame = write_csv()
    write_markdown(frame)
    write_docx(frame)
    print(f"Saved TRIPOD+AI checklist with {len(frame)} items to {OUTPUTS}")


if __name__ == "__main__":
    main()
